import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
from datetime import datetime

# ── Page config ────────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Sistema de Avaliação",
    page_icon="🎓",
    layout="wide",
)

# ── Custom CSS ──────────────────────────────────────────────────────────────────
st.markdown("""
<style>
    .main { background-color: #f8f9fa; }
    .stApp { font-family: Arial, sans-serif; }
    .header-card {
        background: linear-gradient(135deg, #1a237e 0%, #283593 100%);
        padding: 2rem;
        border-radius: 12px;
        color: white;
        text-align: center;
        margin-bottom: 2rem;
        box-shadow: 0 4px 15px rgba(26,35,126,0.3);
    }
    .metric-card {
        background: white;
        padding: 1.2rem;
        border-radius: 10px;
        text-align: center;
        box-shadow: 0 2px 8px rgba(0,0,0,0.08);
        border-top: 4px solid;
    }
    .card-green { border-color: #2e7d32; }
    .card-orange { border-color: #e65100; }
    .card-red { border-color: #c62828; }
    .status-aprovado { color: #2e7d32; font-weight: bold; }
    .status-recuperacao { color: #e65100; font-weight: bold; }
    .status-reprovado { color: #c62828; font-weight: bold; }
    .student-row {
        background: white;
        padding: 0.8rem 1rem;
        border-radius: 8px;
        margin-bottom: 0.4rem;
        box-shadow: 0 1px 4px rgba(0,0,0,0.06);
        display: flex;
        align-items: center;
    }
    div[data-testid="stNumberInput"] input { font-size: 1rem; }
    .section-title {
        font-size: 1.1rem;
        font-weight: 700;
        color: #1a237e;
        border-bottom: 2px solid #e3f2fd;
        padding-bottom: 0.4rem;
        margin-bottom: 1rem;
    }
</style>
""", unsafe_allow_html=True)

# ── Session state init ──────────────────────────────────────────────────────────
if "alunos" not in st.session_state:
    st.session_state.alunos = []
if "sala" not in st.session_state:
    st.session_state.sala = ""
if "etapa" not in st.session_state:
    st.session_state.etapa = "config"  # config | cadastro | resultado
if "idx_atual" not in st.session_state:
    st.session_state.idx_atual = 0
if "total_alunos" not in st.session_state:
    st.session_state.total_alunos = 0

# ── Helper functions ────────────────────────────────────────────────────────────
def get_situacao(media):
    if media >= 7:
        return "Aprovado"
    elif media >= 5:
        return "Recuperação"
    return "Reprovado"

def situacao_html(s):
    cls = {
        "Aprovado": "status-aprovado",
        "Recuperação": "status-recuperacao",
        "Reprovado": "status-reprovado",
    }.get(s, "")
    return f'<span class="{cls}">{s}</span>'

def gerar_xlsx(sala, alunos):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultado"

    # Styles
    azul_escuro = "1A237E"
    azul_claro  = "E3F2FD"
    verde       = "E8F5E9"
    laranja     = "FFF3E0"
    vermelho    = "FFEBEE"
    cinza       = "F5F5F5"

    header_font    = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    subheader_font = Font(name="Arial", bold=True, color=azul_escuro, size=10)
    cell_font      = Font(name="Arial", size=10)
    center         = Alignment(horizontal="center", vertical="center")
    left           = Alignment(horizontal="left", vertical="center")
    thin_border    = Border(
        left=Side(style="thin", color="BDBDBD"),
        right=Side(style="thin", color="BDBDBD"),
        top=Side(style="thin", color="BDBDBD"),
        bottom=Side(style="thin", color="BDBDBD"),
    )

    # Title row
    ws.merge_cells("A1:H1")
    ws["A1"] = f"RELATÓRIO DE AVALIAÇÕES — Sala: {sala}"
    ws["A1"].font = Font(name="Arial", bold=True, color="FFFFFF", size=13)
    ws["A1"].fill = PatternFill("solid", fgColor=azul_escuro)
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = Font(name="Arial", italic=True, color="757575", size=9)
    ws["A2"].alignment = center
    ws["A2"].fill = PatternFill("solid", fgColor=cinza)
    ws.row_dimensions[2].height = 18

    # Column headers
    headers = ["#", "Nome do Aluno", "Nota 1", "Nota 2", "Nota 3", "Média", "Situação", "Observação"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor="283593")
        cell.alignment = center
        cell.border = thin_border
    ws.row_dimensions[3].height = 22

    # Data rows
    cor_situacao = {
        "Aprovado":    verde,
        "Recuperação": laranja,
        "Reprovado":   vermelho,
    }
    for i, a in enumerate(alunos, 1):
        row = i + 3
        cor = cor_situacao.get(a["situacao"], "FFFFFF")
        obs = (
            "Parabéns!" if a["situacao"] == "Aprovado"
            else "Atenção: recuperação necessária" if a["situacao"] == "Recuperação"
            else "Reprovado por média insuficiente"
        )
        valores = [i, a["nome"], a["nota1"], a["nota2"], a["nota3"], "", a["situacao"], obs]
        for col, val in enumerate(valores, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = cell_font
            cell.fill = PatternFill("solid", fgColor=cor)
            cell.border = thin_border
            cell.alignment = center if col != 2 else left
        # Média com fórmula
        media_cell = ws.cell(row=row, column=6)
        media_cell.value = f"=AVERAGE(C{row}:E{row})"
        media_cell.number_format = "0.00"
        media_cell.font = Font(name="Arial", bold=True, size=10)
        media_cell.fill = PatternFill("solid", fgColor=cor)
        media_cell.border = thin_border
        media_cell.alignment = center

    # Summary block
    n = len(alunos)
    sr = n + 5  # start row for summary
    ws.merge_cells(f"A{sr}:H{sr}")
    ws[f"A{sr}"] = "RESUMO DA TURMA"
    ws[f"A{sr}"].font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    ws[f"A{sr}"].fill = PatternFill("solid", fgColor=azul_escuro)
    ws[f"A{sr}"].alignment = center
    ws.row_dimensions[sr].height = 24

    summary = [
        ("Total de Alunos",      n,         "FFFFFF"),
        ("Aprovados",            sum(1 for a in alunos if a["situacao"] == "Aprovado"),    "C8E6C9"),
        ("Em Recuperação",       sum(1 for a in alunos if a["situacao"] == "Recuperação"), "FFE0B2"),
        ("Reprovados",           sum(1 for a in alunos if a["situacao"] == "Reprovado"),   "FFCDD2"),
        ("Média da Turma (AVA1)",f"=AVERAGE(C4:C{3+n})",  azul_claro),
        ("Média da Turma (AVA2)",f"=AVERAGE(D4:D{3+n})",  azul_claro),
        ("Média da Turma (AVA3)",f"=AVERAGE(E4:E{3+n})",  azul_claro),
        ("Média Geral",          f"=AVERAGE(F4:F{3+n})",  azul_claro),
    ]
    for j, (label, val, cor) in enumerate(summary):
        r = sr + 1 + j
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = subheader_font
        lc.fill = PatternFill("solid", fgColor=cinza)
        lc.border = thin_border
        lc.alignment = left
        ws.merge_cells(f"A{r}:D{r}")
        vc = ws.cell(row=r, column=5, value=val)
        vc.font = Font(name="Arial", bold=True, size=10)
        vc.fill = PatternFill("solid", fgColor=cor)
        vc.border = thin_border
        vc.alignment = center
        vc.number_format = "0.00" if "Média" in label else "General"
        ws.merge_cells(f"E{r}:H{r}")
        ws.row_dimensions[r].height = 20

    # Column widths
    widths = [5, 30, 10, 10, 10, 10, 15, 35]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def enviar_email(destinatario, assunto, corpo, xlsx_bytes, nome_arquivo, smtp_user, smtp_pass):
    msg = MIMEMultipart()
    msg["From"]    = smtp_user
    msg["To"]      = destinatario
    msg["Subject"] = assunto
    msg.attach(MIMEText(corpo, "plain", "utf-8"))

    part = MIMEBase("application", "octet-stream")
    part.set_payload(xlsx_bytes.read())
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", f'attachment; filename="{nome_arquivo}"')
    msg.attach(part)

    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
        server.login(smtp_user, smtp_pass)
        server.sendmail(smtp_user, destinatario, msg.as_string())

# ══════════════════════════════════════════════════════════════════════════════
# HEADER
# ══════════════════════════════════════════════════════════════════════════════
st.markdown("""
<div class="header-card">
    <h1 style="margin:0;font-size:2rem;">🎓 Sistema de Avaliação Escolar</h1>
    <p style="margin:0.4rem 0 0;opacity:0.85;">Cadastro de notas, relatório e envio ao professor</p>
</div>
""", unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# ETAPA 1 — Configuração
# ══════════════════════════════════════════════════════════════════════════════
if st.session_state.etapa == "config":
    st.markdown('<div class="section-title">⚙️ Configuração da Turma</div>', unsafe_allow_html=True)
    col1, col2 = st.columns([2, 1])
    with col1:
        sala = st.text_input("Nome / código da sala", placeholder="Ex: 3º Ano A")
    with col2:
        total = st.number_input("Quantidade de alunos", min_value=1, max_value=100, value=5, step=1)

    if st.button("▶ Iniciar Cadastro", type="primary", use_container_width=True):
        if not sala.strip():
            st.error("Informe o nome da sala.")
        else:
            st.session_state.sala        = sala.strip()
            st.session_state.total_alunos = int(total)
            st.session_state.alunos      = []
            st.session_state.idx_atual   = 0
            st.session_state.etapa       = "cadastro"
            st.rerun()

# ══════════════════════════════════════════════════════════════════════════════
# ETAPA 2 — Cadastro dos alunos
# ══════════════════════════════════════════════════════════════════════════════
elif st.session_state.etapa == "cadastro":
    total = st.session_state.total_alunos
    idx   = st.session_state.idx_atual

    # Progress bar
    progress = idx / total
    st.progress(progress, text=f"Aluno {idx + 1} de {total}")

    st.markdown(f'<div class="section-title">👤 Cadastro — Aluno {idx + 1}/{total}</div>', unsafe_allow_html=True)

    with st.form(key=f"form_aluno_{idx}"):
        nome = st.text_input("Nome completo do aluno")
        c1, c2, c3 = st.columns(3)
        with c1:
            n1 = st.number_input("Nota 1 (AVA1)", min_value=0.0, max_value=10.0, step=0.1, format="%.1f")
        with c2:
            n2 = st.number_input("Nota 2 (AVA2)", min_value=0.0, max_value=10.0, step=0.1, format="%.1f")
        with c3:
            n3 = st.number_input("Nota 3 (AVA3)", min_value=0.0, max_value=10.0, step=0.1, format="%.1f")

        submitted = st.form_submit_button("➕ Salvar e próximo", type="primary", use_container_width=True)

        if submitted:
            if not nome.strip():
                st.error("Informe o nome do aluno.")
            else:
                media    = (n1 + n2 + n3) / 3
                situacao = get_situacao(media)
                st.session_state.alunos.append({
                    "nome": nome.strip(), "nota1": n1, "nota2": n2,
                    "nota3": n3, "media": media, "situacao": situacao,
                })
                st.session_state.idx_atual += 1
                if st.session_state.idx_atual >= total:
                    st.session_state.etapa = "resultado"
                st.rerun()

    # Alunos já cadastrados
    if st.session_state.alunos:
        st.markdown("---")
        st.markdown('<div class="section-title">📋 Alunos cadastrados</div>', unsafe_allow_html=True)
        for a in st.session_state.alunos:
            st.markdown(
                f"**{a['nome']}** &nbsp;|&nbsp; "
                f"AVA1: {a['nota1']:.1f} · AVA2: {a['nota2']:.1f} · AVA3: {a['nota3']:.1f} &nbsp;|&nbsp; "
                f"Média: **{a['media']:.2f}** &nbsp;|&nbsp; {situacao_html(a['situacao'])}",
                unsafe_allow_html=True,
            )

    if st.button("🔄 Recomeçar", use_container_width=True):
        st.session_state.etapa = "config"
        st.rerun()

# ══════════════════════════════════════════════════════════════════════════════
# ETAPA 3 — Resultado + Download + E-mail
# ══════════════════════════════════════════════════════════════════════════════
elif st.session_state.etapa == "resultado":
    sala   = st.session_state.sala
    alunos = st.session_state.alunos

    # ── Summary metrics ──────────────────────────────────────────────────────
    aprovados    = sum(1 for a in alunos if a["situacao"] == "Aprovado")
    recuperacao  = sum(1 for a in alunos if a["situacao"] == "Recuperação")
    reprovados   = sum(1 for a in alunos if a["situacao"] == "Reprovado")
    media_geral  = sum(a["media"] for a in alunos) / len(alunos)
    n            = len(alunos)

    st.markdown(f'<div class="section-title">📊 Resultado Final — Sala: {sala}</div>', unsafe_allow_html=True)

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Total de Alunos", n)
    col2.metric("✅ Aprovados",    aprovados,   delta=f"{aprovados/n*100:.0f}%")
    col3.metric("⚠️ Recuperação",  recuperacao, delta=f"{recuperacao/n*100:.0f}%")
    col4.metric("❌ Reprovados",   reprovados,  delta=f"{reprovados/n*100:.0f}%")

    # Médias por avaliação
    st.markdown("---")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Média AVA1", f"{sum(a['nota1'] for a in alunos)/n:.2f}")
    c2.metric("Média AVA2", f"{sum(a['nota2'] for a in alunos)/n:.2f}")
    c3.metric("Média AVA3", f"{sum(a['nota3'] for a in alunos)/n:.2f}")
    c4.metric("Média Geral", f"{media_geral:.2f}")

    # ── Table ────────────────────────────────────────────────────────────────
    st.markdown("---")
    st.markdown('<div class="section-title">📋 Lista de Alunos</div>', unsafe_allow_html=True)

    df = pd.DataFrame(alunos)[["nome", "nota1", "nota2", "nota3", "media", "situacao"]]
    df.columns = ["Nome", "AVA1", "AVA2", "AVA3", "Média", "Situação"]
    df.index   = range(1, len(df) + 1)

    def color_situacao(val):
        colors = {
            "Aprovado":    "background-color:#c8e6c9;color:#1b5e20;font-weight:bold",
            "Recuperação": "background-color:#ffe0b2;color:#bf360c;font-weight:bold",
            "Reprovado":   "background-color:#ffcdd2;color:#b71c1c;font-weight:bold",
        }
        return colors.get(val, "")

    styled = (
        df.style
        .applymap(color_situacao, subset=["Situação"])
        .format({"AVA1": "{:.1f}", "AVA2": "{:.1f}", "AVA3": "{:.1f}", "Média": "{:.2f}"})
        .set_properties(**{"text-align": "center"})
        .set_table_styles([{"selector": "th", "props": [("text-align", "center"), ("background-color", "#283593"), ("color", "white")]}])
    )
    st.dataframe(styled, use_container_width=True, height=min(400, 50 + 36 * n))

    # ── Download ─────────────────────────────────────────────────────────────
    st.markdown("---")
    st.markdown('<div class="section-title">📥 Exportar Planilha</div>', unsafe_allow_html=True)

    xlsx_buf   = gerar_xlsx(sala, alunos)
    nome_arquivo = f"avaliacao_{sala.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    st.download_button(
        label="⬇️ Baixar planilha Excel",
        data=xlsx_buf,
        file_name=nome_arquivo,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )

    # ── E-mail ───────────────────────────────────────────────────────────────
    st.markdown("---")
    st.markdown('<div class="section-title">📧 Enviar por E-mail ao Professor</div>', unsafe_allow_html=True)

    with st.expander("🔧 Configurar e enviar e-mail", expanded=False):
        st.info("ℹ️ Configure o Gmail do **remetente** com uma [Senha de App](https://myaccount.google.com/apppasswords) (não a senha normal). O 2FA precisa estar ativado.")

        col_a, col_b = st.columns(2)
        with col_a:
            email_dest  = st.text_input("E-mail do professor (destinatário)")
            email_rem   = st.text_input("Seu Gmail (remetente)")
        with col_b:
            senha_app   = st.text_input("Senha de App do Gmail", type="password")
            assunto     = st.text_input("Assunto", value=f"Resultado Avaliações — Sala {sala}")

        corpo = st.text_area(
            "Mensagem",
            value=(
                f"Prezado(a) Professor(a),\n\n"
                f"Segue em anexo a planilha com o resultado das avaliações da sala {sala}.\n\n"
                f"Resumo:\n"
                f"  • Aprovados: {aprovados}\n"
                f"  • Em Recuperação: {recuperacao}\n"
                f"  • Reprovados: {reprovados}\n"
                f"  • Média Geral: {media_geral:.2f}\n\n"
                f"Atenciosamente."
            ),
            height=180,
        )

        if st.button("📤 Enviar E-mail", type="primary", use_container_width=True):
            if not all([email_dest, email_rem, senha_app]):
                st.error("Preencha todos os campos de e-mail.")
            else:
                with st.spinner("Enviando..."):
                    try:
                        xlsx_buf2 = gerar_xlsx(sala, alunos)
                        enviar_email(email_dest, assunto, corpo, xlsx_buf2, nome_arquivo, email_rem, senha_app)
                        st.success(f"✅ E-mail enviado com sucesso para **{email_dest}**!")
                    except smtplib.SMTPAuthenticationError:
                        st.error("❌ Falha na autenticação. Verifique o Gmail e a Senha de App.")
                    except Exception as e:
                        st.error(f"❌ Erro ao enviar: {e}")

    # ── Restart ──────────────────────────────────────────────────────────────
    st.markdown("---")
    if st.button("🔄 Nova Avaliação", use_container_width=True):
        for key in ["alunos", "sala", "etapa", "idx_atual", "total_alunos"]:
            del st.session_state[key]
        st.rerun()